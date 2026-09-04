from pathlib import Path

from django.conf import settings
from django.core.management.base import BaseCommand, CommandError
from django.urls import resolve
from openpyxl import load_workbook


class Command(BaseCommand):
    help = "Verifikasi route, template, rich editor, favicon, dan aset integrasi CPNS JUARA v2.5.1."

    def handle(self, *args, **options):
        checks = [
            ("Dashboard", "/"),
            ("Django Admin", "/admin/"),
            ("Content Studio", "/content/"),
            ("Import Excel", "/content/import/"),
            ("Riwayat Import", "/content/import/history/"),
            ("API Health", "/api/health/"),
            ("API Packages", "/api/packages/"),
        ]
        failures = []

        for label, path in checks:
            try:
                match = resolve(path)
                self.stdout.write(
                    self.style.SUCCESS(
                        f"[OK] {label}: {path} -> {match.view_name or match.func.__name__}"
                    )
                )
            except Exception as exc:
                failures.append(f"{label} {path}: {exc}")
                self.stdout.write(self.style.ERROR(f"[FAIL] {label}: {path} -> {exc}"))

        template = (
            Path(settings.BASE_DIR)
            / "static"
            / "templates"
            / "CPNS-JUARA-Template-Import-Soal.xlsx"
        )
        if not template.exists():
            failures.append("Template Excel tidak ditemukan")
            self.stdout.write(
                self.style.ERROR(f"[FAIL] Template Excel tidak ditemukan: {template}")
            )
        else:
            self.stdout.write(self.style.SUCCESS(f"[OK] Template Excel: {template.name}"))
            try:
                wb = load_workbook(template, read_only=True, data_only=True)
                if wb.sheetnames != ["KISI-KISI", "SOAL"]:
                    failures.append(
                        f"Sheet template harus KISI-KISI dan SOAL; ditemukan {wb.sheetnames}"
                    )
                    self.stdout.write(
                        self.style.ERROR(f"[FAIL] Sheet template: {wb.sheetnames}")
                    )
                else:
                    self.stdout.write(
                        self.style.SUCCESS("[OK] Template: KISI-KISI + SOAL dengan stimulus inline")
                    )

                headers = [str(x.value or "").strip().lower() for x in next(wb["SOAL"].iter_rows())]
                for required in [
                    "stimulus_jenis",
                    "stimulus_teks",
                    "stimulus_image_url",
                    "stimulus_sumber",
                    "pertanyaan",
                    "opsi_a", "opsi_b", "opsi_c", "opsi_d", "opsi_e",
                ]:
                    if required not in headers:
                        failures.append(f"Kolom SOAL tidak ditemukan: {required}")
                if not any("Kolom SOAL" in f for f in failures):
                    self.stdout.write(
                        self.style.SUCCESS("[OK] Stimulus tetap terstruktur; opsi A-E masing-masing satu kolom rich content")
                    )
            except Exception as exc:
                failures.append(f"Template Excel gagal diperiksa: {exc}")
                self.stdout.write(self.style.ERROR(f"[FAIL] Template Excel: {exc}"))

        favicon = Path(settings.BASE_DIR) / "static" / "favicon.ico"
        favicon_png = Path(settings.BASE_DIR) / "static" / "favicon-32.png"
        if favicon.exists() and favicon_png.exists():
            self.stdout.write(self.style.SUCCESS("[OK] Favicon CPNS JUARA tersedia"))
        else:
            failures.append("Favicon CPNS JUARA tidak lengkap")
            self.stdout.write(self.style.ERROR("[FAIL] Favicon CPNS JUARA tidak lengkap"))

        rich_editor = Path(settings.BASE_DIR) / "static" / "rich-editor.js"
        if rich_editor.exists():
            editor_source = rich_editor.read_text(encoding="utf-8")
            if "Preview Web + LaTeX" in editor_source or "cpns-rich-preview" in editor_source:
                failures.append("Panel Preview Web + LaTeX masih aktif")
                self.stdout.write(self.style.ERROR("[FAIL] Preview Web + LaTeX masih aktif"))
            else:
                self.stdout.write(self.style.SUCCESS("[OK] Preview Web + LaTeX sudah dihapus"))
        else:
            failures.append("rich-editor.js tidak ditemukan")
            self.stdout.write(self.style.ERROR("[FAIL] rich-editor.js tidak ditemukan"))

        sample_chart = Path(settings.BASE_DIR) / "static" / "sample_stimulus_chart.png"
        if sample_chart.exists():
            self.stdout.write(self.style.SUCCESS("[OK] Contoh grafik stimulus tersedia"))
        else:
            failures.append("Contoh grafik stimulus tidak ditemukan")
            self.stdout.write(self.style.ERROR("[FAIL] Contoh grafik stimulus tidak ditemukan"))

        if failures:
            raise CommandError("Verifikasi gagal: " + " | ".join(failures))

        self.stdout.write(
            self.style.SUCCESS("SEMUA CHECKPOINT CPNS JUARA v2.5.1 LULUS.")
        )
